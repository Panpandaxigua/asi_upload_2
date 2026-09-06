# -*- coding: utf-8 -*-
import tempfile
import unittest
import inspect
from pathlib import Path

from openpyxl import Workbook, load_workbook

from asi_bot import AsiBot
from tools import (
    iter_image_files, read_excel, run_precheck, split_multi_value,
    write_result_to_excel,
)
from upload_controls import default_upload_controls


SHEET_NAME = "ASI制作表格"


def make_workbook(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for row in range(1, 5):
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header)
    for row_offset, values in enumerate(rows, 5):
        for col, value in enumerate(values, 1):
            ws.cell(row=row_offset, column=col, value=value)
    wb.save(path)
    wb.close()


class ExcelProcessingTests(unittest.TestCase):
    def test_read_excel_returns_blank_and_non_success_process_rows_and_keeps_excel_index(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "products.xlsx"
            make_workbook(
                path,
                ["Product_Number", "Product_Name", "Process", "Process status"],
                [
                    ["SKU-1", "Ready", "", ""],
                    ["SKU-2", "Done", "success", ""],
                    ["SKU-3", "Retry", "fail", "Make Active failed"],
                    ["SKU-4", "Problem", "problem", "Validation warning"],
                    ["", "No number", "", ""],
                ],
            )

            df = read_excel(path)

            self.assertEqual(df["Product_Number"].tolist(), ["SKU-1", "SKU-3", "SKU-4"])
            self.assertEqual(df.index.tolist(), [0, 2, 3])

    def test_write_result_uses_process_for_status_and_process_status_for_message(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "products.xlsx"
            make_workbook(
                path,
                ["Product_Number", "Product_Name", "Process", "Process status"],
                [["SKU-1", "Ready", "", ""]],
            )

            self.assertTrue(write_result_to_excel(path, 0, "fail", "Make Active validation failed"))

            wb = load_workbook(path)
            ws = wb[SHEET_NAME]
            headers = [cell.value for cell in ws[4]]
            process_col = headers.index("Process") + 1
            process_status_col = headers.index("Process status") + 1
            self.assertEqual(ws.cell(row=5, column=process_col).value, "fail")
            self.assertEqual(ws.cell(row=5, column=process_status_col).value, "Make Active validation failed")
            wb.close()

    def test_write_result_migrates_legacy_process_status_to_process(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "products.xlsx"
            make_workbook(
                path,
                ["Product_Number", "Product_Name", "Process Status"],
                [["SKU-1", "Ready", ""]],
            )

            self.assertTrue(write_result_to_excel(path, 0, "success", ""))

            wb = load_workbook(path)
            ws = wb[SHEET_NAME]
            headers = [cell.value for cell in ws[4]]
            process_col = headers.index("Process") + 1
            process_status_col = headers.index("Process status") + 1
            self.assertEqual(ws.cell(row=5, column=process_col).value, "success")
            self.assertEqual(ws.cell(row=5, column=process_status_col).value, None)
            wb.close()

    def test_precheck_allows_custom_material_names(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "products.xlsx"
            make_workbook(
                path,
                [
                    "Product_Number",
                    "Product_Name",
                    "Material",
                    "Material Custom Name",
                ],
                [["SKU-1", "Custom product", "My Custom Composite", "Eco Composite"]],
            )

            result = run_precheck(path)

            material_issues = [
                issue
                for warning in result.get("warnings", [])
                for issue in warning.get("issues", [])
                if "Material" in issue
            ]
            self.assertEqual(material_issues, [])

    def test_precheck_uses_required_upload_controls(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "products.xlsx"
            make_workbook(
                path,
                [
                    "Product_Number",
                    "Product_Name",
                    "Q1",
                    "P1",
                    "Process",
                    "Process status",
                ],
                [["SKU-1", "Missing price", "100", "", "", ""]],
            )
            controls = default_upload_controls()
            controls["modules"]["pricing"]["fields"]["price_tiers"] = "required"

            result = run_precheck(path, upload_controls=controls)

            self.assertEqual(result["problems"], 1)
            issues = result["warnings"][0]["issues"]
            self.assertTrue(any("价格档位" in issue for issue in issues))
            wb = load_workbook(path)
            ws = wb[SHEET_NAME]
            self.assertEqual(ws["E5"].value, "problem")
            self.assertIn("价格档位", ws["F5"].value)
            wb.close()

    def test_split_multi_value_accepts_common_chinese_punctuation(self):
        text = "Red\uff0cBlue\uff1bGreen\u3001Yellow\uff0cBlack"
        self.assertEqual(split_multi_value(text), ["Red", "Blue", "Green", "Yellow", "Black"])

    def test_artwork_information_accepts_chinese_punctuation_comments(self):
        text = "1\uff1afree\uff1b2\uff0c3\uff1aNeed vector art;4\uff1afree"
        items = AsiBot.parse_artwork_information_items(text)
        self.assertEqual(
            items,
            [
                {"name": "Virtual Proof", "comment": "free"},
                {"name": "Paper Proof", "comment": ""},
                {"name": "Art Services", "comment": "Need vector art"},
                {"name": "Pre-production Proof", "comment": "free"},
            ],
        )

    def test_key_value_size_accepts_colon_or_equals(self):
        text = "L\uff1a1 in\uff1bW=2 in;H:3 in"
        self.assertEqual(
            AsiBot.parse_key_value_size(text),
            [("L", "1 in"), ("W", "2 in"), ("H", "3 in")],
        )

    def test_volume_and_capacity_sizes_update_criteria_observables(self):
        source = inspect.getsource(AsiBot.fill_product_size)

        self.assertIn("criteriaValueObj.Value", source)
        self.assertIn("criteriaValueObj.UOM", source)
        self.assertNotIn("setObs(hold, 'value'", source)

    def test_setup_charges_accept_chinese_punctuation(self):
        text = "Front\uff1a50\uff1bCenter:30"
        charges = AsiBot.parse_setup_charges(text, ["Front", "Center"])
        self.assertEqual(
            charges,
            [
                {"location": "Front", "amount": "50", "code": ""},
                {"location": "Center", "amount": "30", "code": ""},
            ],
        )

    def test_setup_charges_are_keyed_by_imprint_location_with_separate_code(self):
        payload = AsiBot.build_pricing_payload({
            "Imprint Location": "Front, Center",
            "Set-up Change": "Front:35, Center:36",
            "Set-up Change Codes": "C/R 40%",
            "Price Codes": "A/P 50%",
        })

        self.assertEqual(payload["setupChargeCode"], "C/R 40%")
        self.assertEqual(payload["setupChargeCodeValue"], "R")
        self.assertEqual(
            payload["setupCharges"],
            [
                {"location": "Front", "amount": "35", "code": "R"},
                {"location": "Center", "amount": "36", "code": "R"},
            ],
        )

    def test_pricing_payload_supports_eight_tiers(self):
        row = {f"Q{i}": i * 100 for i in range(1, 9)}
        row.update({f"P{i}": 20 - i for i in range(1, 9)})

        payload = AsiBot.build_pricing_payload(row)

        self.assertEqual(len(payload["prices"]), 8)
        self.assertEqual(
            payload["prices"][-1],
            {"quantity": "800", "netCost": "12"},
        )

    def test_material_custom_names_align_with_materials(self):
        entries = AsiBot.build_material_entries({
            "Material": "Polyester, Wood",
            "Material Custom Name": "Recycled Polyester, Bamboo Finish",
        })

        self.assertEqual(
            entries,
            [
                {"material": "Polyester", "alias": "Recycled Polyester"},
                {"material": "Wood", "alias": "Bamboo Finish"},
            ],
        )

    def test_unmatched_material_defaults_alias_to_requested_name(self):
        entries = AsiBot.build_material_entries({
            "Material": "My Custom Composite",
        })

        self.assertEqual(
            entries,
            [{"material": "My Custom Composite", "alias": "My Custom Composite"}],
        )

    def test_material_selector_passes_alias_to_page_model(self):
        calls = []

        class FakeTab:
            def run_js(self, script, *args):
                calls.append((script, args))
                return {
                    "ok": True,
                    "wanted": args[0],
                    "alias": args[1],
                    "matched": "Polyester",
                    "selected": args[1],
                    "fallback": False,
                }

        bot = AsiBot.__new__(AsiBot)
        bot.tab = FakeTab()

        result = bot.select_material_type([
            {"material": "Polyester", "alias": "Recycled Polyester"},
        ])

        self.assertTrue(result["ok"])
        self.assertEqual(calls[0][1], ("Polyester", "Recycled Polyester", 0))
        self.assertIn("aliasName", calls[0][0])
        self.assertIn("toggleMaterialTypeTemplate", calls[0][0])
        self.assertIn("Other", calls[0][0])

    def test_setup_charge_locations_extend_imprint_locations(self):
        payload = AsiBot.build_imprint_payload({
            "Imprint Location": "Front",
            "Set-up Change": "Front:35, Center:36",
        })

        self.assertEqual(payload["imprintLocations"], ["Front", "Center"])

    def test_setup_charges_parse_locations_even_when_imprint_location_column_is_partial(self):
        payload = AsiBot.build_pricing_payload({
            "Imprint Location": "Front",
            "Set-up Change": "Front:35, Center:36",
            "Set-up Change Codes": "C/R 40%",
        })

        self.assertEqual(
            payload["setupCharges"],
            [
                {"location": "Front", "amount": "35", "code": "R"},
                {"location": "Center", "amount": "36", "code": "R"},
            ],
        )

    def test_setup_charge_codes_can_match_locations_individually(self):
        charges = AsiBot.parse_setup_charges(
            "Front:35; Center:36",
            ["Front", "Center"],
            "Front:C; Center:R",
        )

        self.assertEqual(
            charges,
            [
                {"location": "Front", "amount": "35", "code": "C"},
                {"location": "Center", "amount": "36", "code": "R"},
            ],
        )

    def test_media_tags_pass_image_stems_for_dynamic_color_matching(self):
        tags = AsiBot.media_tags_for_files([
            r"C:\images\1_burgundy+gold.png",
            r"C:\images\Product Hero.JPG",
        ])

        self.assertEqual(
            tags,
            [
                {"stem": "1_burgundy+gold"},
                {"stem": "Product Hero"},
            ],
        )

    def test_media_tags_include_colors_from_available_product_colors(self):
        tags = AsiBot.media_tags_for_files(
            [r"C:\images\2_Orange_Navy Blue.png"],
            ["Orange", "Navy Blue", "Black"],
        )

        self.assertEqual(tags, [{"stem": "2_Orange_Navy Blue", "colors": ["Orange", "Navy Blue"]}])

    def test_media_color_extraction_uses_available_product_colors(self):
        self.assertEqual(
            AsiBot.media_colors_from_name(
                "shirt_red_yellow.png",
                ["Black (Black)", "Blue (Blue)", "Burgundy (Red)", "Gold (Yellow)", "Gray (Gray)", "Green (Green)"],
            ),
            ["Burgundy (Red)", "Gold (Yellow)"],
        )

    def test_media_color_extraction_prefers_longer_available_color_names(self):
        self.assertEqual(
            AsiBot.media_colors_from_name("bag_navy-blue.png", ["Navy Blue (Blue)", "Blue (Blue)"]),
            ["Navy Blue (Blue)"],
        )

    def test_media_color_extraction_ignores_leading_image_order_number(self):
        self.assertEqual(
            AsiBot.media_colors_from_name(
                "1_White_navy blue.png",
                ["White (White)", "Navy Blue (Blue)", "Blue (Blue)"],
            ),
            ["White (White)", "Navy Blue (Blue)"],
        )

    def test_media_color_extraction_uses_name_when_no_parenthetical_color_exists(self):
        self.assertEqual(
            AsiBot.media_colors_from_name("bag_burgundy.png", ["Burgundy", "Gold"]),
            ["Burgundy"],
        )

    def test_iter_image_files_uses_leading_number_as_upload_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            for name in ["10_green.png", "2_black.png", "1_white.png", "no_order.png"]:
                (folder / name).write_bytes(b"image")

            ordered = [Path(path).name for path in iter_image_files(folder)]

            self.assertEqual(ordered, ["1_white.png", "2_black.png", "10_green.png", "no_order.png"])

    def test_product_colors_are_canonicalized_and_deduped(self):
        colors = AsiBot.parse_product_colors("White, navy blue, pink, black, Navy Blue, BLUE")
        self.assertEqual(colors, ["White", "Navy Blue", "Pink", "Black", "Blue"])

    def test_blank_personalization_available_stays_empty_in_imprint_payload(self):
        payload = AsiBot.build_imprint_payload({
            "Select Method": "Silkscreen",
            "Personalization Available": "",
            "Artwork Information": "1",
        })
        self.assertFalse(payload["hasPersonalizationAvailable"])
        self.assertEqual(payload["personalizationMethods"], [])

    def test_imprint_colors_column_name_matches_loose_spacing(self):
        payload = AsiBot.build_imprint_payload({
            " imprint colors ": "1, 2, 3",
        })
        self.assertTrue(payload["hasImprintColorCodes"])
        self.assertEqual(payload["imprintColors"], {"standard": True, "custom": True, "pms": True})

    def test_imprint_location_supports_multiple_values(self):
        payload = AsiBot.build_imprint_payload({
            "Imprint Location": "Front, Center\uff1bBack",
        })
        self.assertEqual(payload["imprintLocations"], ["Front", "Center", "Back"])


if __name__ == "__main__":
    unittest.main()
